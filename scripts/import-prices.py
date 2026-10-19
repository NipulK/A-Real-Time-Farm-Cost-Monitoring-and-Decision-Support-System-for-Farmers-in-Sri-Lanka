"""Read only the two candidate workbooks. Never combine overlapping versions."""
import json,hashlib,datetime,argparse
from pathlib import Path
import openpyxl
parser=argparse.ArgumentParser();parser.add_argument('reference',type=Path);args=parser.parse_args()
root=Path(__file__).resolve().parents[1]; records=[]; sources=[]; products=[]
for market,filename in [('Dambulla','dambulla_price_comparison_bilingual_categorized.xlsx'),('Thabuththegama','thabuththegama_price_comparison_bilingual_categorized.xlsx')]:
 f=args.reference/'work_price_list/outputs/categorized-products'/filename
 sha=hashlib.sha256(f.read_bytes()).hexdigest();w=openpyxl.load_workbook(f,read_only=True,data_only=True)
 for s in w:
  rows=list(s.iter_rows(values_only=True));dates=[(i,v.date().isoformat()) for i,v in enumerate(rows[2]) if isinstance(v,datetime.datetime)]
  count=0
  for r,row in enumerate(rows[4:],5):
   label=row[0] if row else None
   if not isinstance(label,str) or not ' - ' in label:continue
   count+=1; category=str(row[len(dates)*2+1] or 'Unclassified') if len(row)>len(dates)*2+1 else 'Unclassified'
   source_category=category
   if any(term in label.lower() for term in ['dry chil','ginger','turmeric']): category='Spices'
   if any(term in label.lower() for term in ['lentil','dhal','gram','cowpea','finger millet','soya bean','sesame','peanuts','corns']): category='Grains & pulses'
   english=label.split(' - ',1)[1]
   canonical=next((name for term,name in [('watermelon','Watermelon'),('guava','Guava'),('carrot','Carrot'),('potato','Potato'),('leeks','Leeks'),('grape','Grapes')] if term in english.lower()),english)
   products.append({'canonicalProduct':canonical,'sourceCategory':source_category,'market':market,'label':label,'category':category,'source':filename,'row':r})
   for col,date in dates:
    lo=row[col] if col<len(row) else None;hi=row[col+1] if col+1<len(row) else None
    lo=lo if isinstance(lo,(int,float)) else None;hi=hi if isinstance(hi,(int,float)) else None
    if lo is None and hi is None:continue
    uncertain=market=='Dambulla' and date>='2026-09-24'
    records.append({'canonicalProduct':canonical,'sourceCategory':source_category,'id':f'{market}-{r}-{date}','product':label.split(' - ',1)[1],'originalLabel':label,'category':category,'originVariant':label.split(' - ',1)[1],'date':date,'lower':lo,'upper':hi,'currency':'LKR','unit':'kg','unitVerification':'Workbook notes only; source verification pending','priceType':'Historical price-list range','claimedMarket':market,'market':'Uncertain (Dambulla / Keppetipola)' if uncertain else market,'marketIdentity':'uncertain' if uncertain else 'claimed','state':'reversed' if lo is not None and hi is not None and lo>hi else 'partial' if lo is None or hi is None else 'unverified','source':filename,'sourceHash':sha,'sheet':s.title,'cells':f'{openpyxl.utils.get_column_letter(col+1)}{r}:{openpyxl.utils.get_column_letter(col+2)}{r}','verification':'unverified'})
  sources.append({'filename':filename,'sha256':sha,'sheet':s.title,'products':count,'dates':[d for _,d in dates]})
(root/'data/prices.json').write_text(json.dumps(records,ensure_ascii=False),encoding='utf8')
(root/'data/price-products.json').write_text(json.dumps(products,ensure_ascii=False),encoding='utf8')
summary={'sources':sources,'observations':len(records),'complete':sum(r['state']=='unverified' for r in records),'partial':sum(r['state']=='partial' for r in records),'reversed':sum(r['state']=='reversed' for r in records),'uncertainMarket':sum(r['marketIdentity']=='uncertain' for r in records),'allImportedUnverified':True}
(root/'data/price-import-audit.json').write_text(json.dumps(summary,indent=2),encoding='utf8');print(json.dumps(summary,indent=2))
